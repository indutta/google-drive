#!/usr/bin/env python3
"""Parse INS_Inspection_Flow.xlsx into plan.js (window.PLAN) for the inspection app.

Usage: python3 -I tools/build_plan.py [source.xlsx] [plan.js]
"""
import json, re, sys, os
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..', 'source', 'INS_Inspection_Flow.xlsx')
OUT = sys.argv[2] if len(sys.argv) > 2 else os.path.join(HERE, '..', 'plan.js')

# Where each step is recorded.  job = once per work order (pattern/tooling),
# log = sand & calibration log (date + shift), heat = once per heat, casting = every casting.
SCOPE = {1: 'job', 2: 'log', 3: 'casting', 4: 'heat', 5: 'casting', 6: 'casting',
         7: 'casting', 8: 'heat', 9: 'casting', 10: 'casting'}
# Stages / steps that depend on the PO / customer QAP (switchable per job)
OPTIONAL_STEPS = {'7.1', '7.2', '7.3', '7.4', '8.1', '8.2', '8.3', '8.4', '9.1'}
# Steps that are sampled (not every casting / heat) -> recorded when sampled, never block the flow
SAMPLED = {'3.2', '8.3'}
# Log frequency classes for the sand / calibration log
LOG_FREQ = [(r'once in three months', 'quarterly'), (r'once in a week', 'weekly'),
            (r'once a day', 'daily'), (r'shift', 'shift')]
# Multi-parameter checks (parameters + limits supplied on the heat record)
CHEM = {'4.1', '4.2'}
MECH = {'8.1'}
# "Mix 2.5 mm" in the source is read as a maximum penetration (typo for "Max") - flagged as assumed
DIM_STEPS = {'1.1', '1.3', '1.4', '5.2', '9.1'}
YN_STEPS = {'8.2'}
ASSUMED = {'2.10.3', '2.11.4', '2.12.4'}

NUM = r'(\d+(?:\.\d+)?)'

def parse_criteria(c):
    """Return dict(kind,min,max,unit) from the acceptance-criteria text."""
    t = (c or '').strip()
    m = re.match(rf'^{NUM}\s*(%?)\s*Max$', t, re.I)
    if m: return dict(kind='max', max=float(m[1]), unit=m[2] or '')
    m = re.match(rf'^{NUM}\s*(kg/cm²)\s*min', t, re.I)
    if m: return dict(kind='min', min=float(m[1]), unit=m[2])
    m = re.match(rf'^Mix {NUM} mm', t)
    if m: return dict(kind='max', max=float(m[1]), unit='mm')
    m = re.match(rf'^{NUM}\s*cc/gm of sand Max', t)
    if m: return dict(kind='max', max=float(m[1]), unit='cc/gm')
    m = re.match(rf'^{NUM}\s*[–-]\s*{NUM}\s*%', t)
    if m: return dict(kind='range', min=float(m[1]), max=float(m[2]), unit='%')
    m = re.match(rf'^{NUM}\s*(?:to|–|-)\s*{NUM}\s*kg', t)
    if m: return dict(kind='range', min=float(m[1]), max=float(m[2]), unit='kg / 15 sec')
    m = re.match(rf'^{NUM}[–-]{NUM}\s*kg of total', t)
    if m: return dict(kind='range', min=float(m[1]), max=float(m[2]), unit='kg / 15 sec')
    m = re.match(rf'^{NUM}[–-]{NUM}\s*Baume', t)
    if m: return dict(kind='range', min=float(m[1]), max=float(m[2]), unit='Baumé')
    m = re.match(rf'^Hardness {NUM} min', t)
    if m: return dict(kind='min', min=float(m[1]), unit='')
    return None

def split_actions(g):
    parts = [p.strip() for p in re.split(r'\s*/\s*', g or '') if p.strip()]
    return parts or ['Other']

def classify(action):
    return 'reject' if re.search(r'reject|destroy|dispos|remelt|scrap|pig the melt', action, re.I) else 'rework'

def log_freq(sample):
    for rx, name in LOG_FREQ:
        if re.search(rx, sample or '', re.I): return name
    return 'shift'

def main():
    wb = openpyxl.load_workbook(SRC)
    ws = wb['Inspection Plan']
    stages, stage, group, step = [], None, None, None
    counters = {}
    for r in range(6, ws.max_row + 1):
        a, b, c, d, e, f, g = [ws.cell(r, i).value for i in range(1, 8)]
        a = str(a).strip() if a is not None else None
        if a and a.startswith('STAGE'):
            m = re.match(r'STAGE (\d+):\s*(.*)', a)
            stage = dict(no=int(m[1]), name=m[2].title().replace('&', '&').replace('(Ndt)', '(NDT)').replace('(If Applicable)', '(if applicable)'),
                         scope=SCOPE[int(m[1])], groups=[])
            stages.append(stage); group = None; continue
        if a and a.startswith('- -'): break
        if stage is None: continue
        if a and not re.match(r'^\d+\.\d+$', a) and not any([b, c, d]):
            group = dict(name=a, steps=[]); stage['groups'].append(group); continue
        if a and re.match(r'^\d+\.\d+$', a):
            if group is None or stage['groups'] == [] :
                group = dict(name=None, steps=[]); stage['groups'].append(group)
            step = dict(id=a, name=b, checks=[], optional=a in OPTIONAL_STEPS, sampled=a in SAMPLED)
            group['steps'].append(step); counters[a] = 0
        if step is None or (d is None and c is None): continue
        counters[step['id']] += 1
        cid = f"{step['id']}.{counters[step['id']]}"
        fills = {col: (ws.cell(r, col).fill.fgColor.rgb if ws.cell(r, col).fill.fill_type else None) for col in (6, 7)}
        tbc = any(v == '00FFF2A8' for v in fills.values())
        chk = dict(id=cid, step=step['id'], sample=c if d else None, param=d, method=e, criteria=f, action=g,
                   actions=split_actions(g) if g else [], tbc=tbc)
        if d is None:                                   # 2.1: whole-row reference to the incoming-material QAP
            chk.update(sample='Per incoming QAP', param='Incoming sand verified', method='Doc No. DCL/QAP/ICI',
                       criteria=c, action='Reject lot / return to supplier', actions=['Reject lot / return to supplier'])
            chk['kind'] = 'yn'
        else:
            p = parse_criteria(f)
            if step['id'] in CHEM: chk['kind'] = 'chem'
            elif step['id'] in MECH: chk['kind'] = 'mech'
            elif p: chk.update(p)
            elif re.search(r'visual|closing|matching|heat cycle|packing|surface|crack|internal|sub-surface|coating|d\.p\.|vent|1 pc', (d or '') + ' ' + (e or ''), re.I) \
                    and not re.search(r'dimension|temperature|hardness|strength|microstructure', d or '', re.I):
                chk['kind'] = 'yn'
            else:
                chk['kind'] = 'spec'                    # numeric reading with limits entered at inspection time
        if 'kind' not in chk: chk['kind'] = 'spec'
        if chk['id'] in ('2.4.1', '2.5.1'): chk['unit'] = '% of sand wt' if chk['id'] == '2.4.1' else '% of resin wt'
        if chk['id'] in ('2.2.3', '2.3.3'): chk['unit'] = '%'
        if step['id'] in DIM_STEPS and re.search('dimension', chk['param'] or '', re.I): chk['kind'] = 'dim'
        if step['id'] in YN_STEPS: chk['kind'] = 'yn'
        if step['id'] == '3.2': chk['unit'] = 'scratch hardness'
        chk['actions'] = [x[:1].upper() + x[1:] for x in chk['actions']]
        if chk['actions'] == ['To be confirmed']:
            chk['actions'] = ['Salvaging', 'Reject']; chk['assumed'] = True
        if cid in ASSUMED: chk['assumed'] = True
        if stage['scope'] == 'log': chk['freq'] = log_freq(c)
        if step['id'] in ('4.3', '4.4'): chk['unit'] = '°C'
        if step['id'] == '8.3': chk['unit'] = 'HBW'
        if step['id'] == '8.4': chk['unit'] = 'J'
        for k in ('min', 'max'):
            pass
        for k, v in list(chk.items()):
            if v is None: del chk[k]
        step['checks'].append(chk)
        if cid == '10.1.1':                            # split visual and dimensional checks
            chk.update(param='Visual inspection', method='Visual', kind='yn')
            c2 = dict(chk, id='10.1.2', param='Specified dimensions as per drawing & spec.', method='Inspection gauge, measuring instruments', kind='dim')
            step['checks'].append(c2)
    # stage level applicability notes
    for s in stages:
        s['optional'] = s['no'] in (7, 8, 9)
    plan = dict(company='DATRE CORPORATION LIMITED', title='Casting Process Flow – Inspection Plan', stages=stages)
    with open(OUT, 'w') as fh:
        fh.write('/* Generated by tools/build_plan.py from source/INS_Inspection_Flow.xlsx - do not edit by hand */\n')
        fh.write('window.PLAN = ' + json.dumps(plan, ensure_ascii=False, indent=1) + ';\n')
    n = sum(len(st['checks']) for s in stages for g in s['groups'] for st in g['steps'])
    print(f'{len(stages)} stages, {sum(len(g["steps"]) for s in stages for g in s["groups"])} steps, {n} checks -> {OUT}')

main()
